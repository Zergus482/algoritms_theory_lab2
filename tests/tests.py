import pytest
from docx import Document
from openpyxl import load_workbook

from vehicles import Vehicle
from car_types import PassengerCar, Truck, Bus
from reports import TripReport


# 1. Абстрактный класс нельзя создать напрямую
def test_vehicle_is_abstract():
    with pytest.raises(TypeError):
        Vehicle("X", 10.0, 1.0)


# 2. Managed-атрибуты отклоняют некорректные значения
def test_managed_attributes_validation():
    car = PassengerCar()
    with pytest.raises(ValueError):
        car.base_consumption = -1
    with pytest.raises(ValueError):
        car.fuel_price = 0
    car.fuel_price = 55
    assert car.fuel_price == 55


# 3. Легковой: расход растёт с загрузкой
def test_passenger_car_consumption_grows_with_load():
    car = PassengerCar(base_consumption=7.0, load_capacity=0.4)
    assert car.consumption(0) == pytest.approx(7.0)
    assert car.consumption(0.2) > car.consumption(0)


# 4. Грузовой: расход считается по формуле
def test_truck_consumption_formula():
    t = Truck(base_consumption=25.0, load_capacity=20.0)
    # 25 * (1 + 0.015 * 10) = 28.75
    assert t.consumption(10) == pytest.approx(28.75)


# 5. Автобус: скорость падает с числом пассажиров
def test_bus_speed_decreases_with_load():
    b = Bus(load_capacity=50.0)
    assert b.avg_speed(0) > b.avg_speed(25) > b.avg_speed(50)


# 6. Проверка диапазона загрузки
def test_load_out_of_range_raises():
    with pytest.raises(ValueError):
        Truck(load_capacity=20.0).consumption(21)
    with pytest.raises(ValueError):
        Bus(load_capacity=50.0).consumption(-1)


# 7. Полный расчёт поездки (топливо, стоимость, время)
def test_trip_calculation():
    car = PassengerCar(base_consumption=10.0, load_capacity=1.0, fuel_price=50.0)
    assert car.fuel_needed(100, 0) == pytest.approx(10.0)   # 10 л
    assert car.trip_cost(100, 0) == pytest.approx(500.0)    # 10 * 50
    assert car.trip_time(90, 0) == pytest.approx(1.0)       # 90 / 90


# 8. Dunder-методы __str__, __repr__, __eq__
def test_dunder_methods():
    a = Truck(base_consumption=25.0, load_capacity=20.0)
    b = Truck(base_consumption=25.0, load_capacity=20.0)
    assert a == b
    assert "Truck" in repr(a)



# 9. TripReport: __len__ и __iter__
def test_report_len_and_iter():
    rep = TripReport()
    rep.add(PassengerCar(), 100, 0.2)
    rep.add(Truck(), 500, 10)
    assert len(rep) == 2
    assert len(list(rep)) == 2


# 10. Экспорт отчёта в .docx и .xlsx
def test_report_export_docx_xlsx(tmp_path):
    rep = TripReport()
    rep.add(PassengerCar(), 100, 0.2)
    rep.add(Truck(), 500, 10)

    docx_path = tmp_path / "r.docx"
    xlsx_path = tmp_path / "r.xlsx"
    rep.to_docx(str(docx_path))
    rep.to_xlsx(str(xlsx_path))

    assert docx_path.exists() and docx_path.stat().st_size > 0
    assert xlsx_path.exists() and xlsx_path.stat().st_size > 0

    doc = Document(str(docx_path))
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 3

    ws = load_workbook(str(xlsx_path)).active
    assert ws.max_row == 3
    assert ws.cell(row=1, column=1).value == "Автомобиль"